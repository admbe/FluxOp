"""Planning assumptions: manual savings/cost line items for the outlook.

The docx feedback that motivated this feature (2026-08-10): the planner
cannot direct the savings estimate from inside the app when part of the
saving depends on manual project work. These tests pin the contract: items
round-trip, enabled items move the projection, disabled items do not, and
the executive workbook lists exactly the items marked for reports.
"""
from datetime import date, datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook
from io import BytesIO

from api.database import FluxDatabase


def make_database(tmp: str) -> FluxDatabase:
    database = FluxDatabase(Path(tmp) / "flux.duckdb")
    database.init()
    return database


def seed_month(database: FluxDatabase, month: date, amount: float) -> None:
    with database.connect() as db:
        db.execute(
            """INSERT INTO monthly_cost_history
               (snapshot_id, observed_at, month, cost_type,
                subscription_id, amount, currency, source)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            ["seed", datetime.now(timezone.utc), month,
             "AmortizedCost", "sub-a", amount, "USD", "test"],
        )
        db.commit()


class PlanningAssumptionsCrudTests(unittest.TestCase):
    def test_round_trip_and_ordering(self):
        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            saved = database.save_planning_assumptions(
                [
                    {
                        "label": "Decommission wave 2",
                        "monthlyAmount": 1200.0,
                        "direction": "saving",
                        "includeInReports": True,
                        "enabled": True,
                        "notes": "12 VMs leaving after the ERP cutover",
                    },
                    {
                        "label": "New security tooling",
                        "monthlyAmount": 300.0,
                        "direction": "cost",
                        "includeInReports": False,
                        "enabled": True,
                        "notes": "",
                    },
                ],
                updated_by="adam",
            )
            self.assertEqual(len(saved), 2)
            self.assertEqual(saved[0]["label"], "Decommission wave 2")
            self.assertEqual(saved[0]["direction"], "saving")
            self.assertTrue(saved[0]["id"])
            self.assertEqual(saved[1]["direction"], "cost")
            self.assertFalse(saved[1]["includeInReports"])
            self.assertEqual(saved[0]["updatedBy"], "adam")

            # net = +1200 saving - 300 cost
            self.assertAlmostEqual(
                database.planning_assumptions_monthly_net(), 900.0
            )

    def test_disabled_items_do_not_count(self):
        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            database.save_planning_assumptions(
                [
                    {
                        "label": "Paused project",
                        "monthlyAmount": 5000.0,
                        "direction": "saving",
                        "enabled": False,
                    }
                ]
            )
            self.assertEqual(database.planning_assumptions_monthly_net(), 0.0)

    def test_save_replaces_wholesale(self):
        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            database.save_planning_assumptions(
                [{"label": "First", "monthlyAmount": 10.0}]
            )
            database.save_planning_assumptions(
                [{"label": "Second", "monthlyAmount": 20.0}]
            )
            items = database.planning_assumptions()
            self.assertEqual([item["label"] for item in items], ["Second"])


class OutlookIntegrationTests(unittest.TestCase):
    def seeded_database(self, tmp: str) -> FluxDatabase:
        database = make_database(tmp)
        for month_index in range(1, 8):
            seed_month(database, date(2026, month_index, 1), 1000.0)
        return database

    def test_enabled_assumption_lowers_projection(self):
        with TemporaryDirectory() as tmp:
            database = self.seeded_database(tmp)
            baseline = database.fiscal_year_outlook(as_of=date(2026, 8, 2))
            database.save_planning_assumptions(
                [
                    {
                        "label": "Decommission",
                        "monthlyAmount": 200.0,
                        "direction": "saving",
                        "enabled": True,
                    }
                ]
            )
            adjusted = database.fiscal_year_outlook(as_of=date(2026, 8, 2))
            self.assertLess(adjusted["fyTotal"], baseline["fyTotal"])
            self.assertTrue(adjusted["manualAssumptions"]["applied"])
            self.assertEqual(
                adjusted["manualAssumptions"]["monthlyNet"], 200.0
            )
            self.assertTrue(
                any(
                    "manually entered planning assumption" in item
                    for item in adjusted["limitations"]
                )
            )

    def test_disabled_assumption_changes_nothing(self):
        with TemporaryDirectory() as tmp:
            database = self.seeded_database(tmp)
            baseline = database.fiscal_year_outlook(as_of=date(2026, 8, 2))
            database.save_planning_assumptions(
                [
                    {
                        "label": "Paused",
                        "monthlyAmount": 200.0,
                        "direction": "saving",
                        "enabled": False,
                    }
                ]
            )
            adjusted = database.fiscal_year_outlook(as_of=date(2026, 8, 2))
            self.assertEqual(adjusted["fyTotal"], baseline["fyTotal"])
            self.assertFalse(adjusted["manualAssumptions"]["applied"])


class WorkbookExportTests(unittest.TestCase):
    def test_export_lists_only_report_marked_items(self):
        from api.executive_export import build_executive_workbook

        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            seed_month(database, date(2026, 7, 1), 1000.0)
            database.save_planning_assumptions(
                [
                    {
                        "label": "Listed item",
                        "monthlyAmount": 100.0,
                        "direction": "saving",
                        "includeInReports": True,
                    },
                    {
                        "label": "Private item",
                        "monthlyAmount": 50.0,
                        "direction": "cost",
                        "includeInReports": False,
                    },
                ]
            )
            payload = build_executive_workbook(database)
            workbook = load_workbook(BytesIO(payload))
            sheet = workbook["Assumptions"]
            cells = [
                str(cell.value)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            self.assertIn("Listed item", cells)
            self.assertNotIn("Private item", cells)

    def test_every_chart_has_visible_axes_and_titles(self):
        """The 2026-08-10 feedback: exported charts had no axis bearings."""
        from api.executive_export import build_executive_workbook

        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            seed_month(database, date(2026, 7, 1), 1000.0)
            payload = build_executive_workbook(database)
            workbook = load_workbook(BytesIO(payload))
            charts = [
                chart
                for sheet in workbook.worksheets
                for chart in sheet._charts
            ]
            self.assertGreaterEqual(len(charts), 1)
            for chart in charts:
                self.assertIs(
                    chart.x_axis.delete, False,
                    "category axis must not be hidden",
                )
                self.assertIs(
                    chart.y_axis.delete, False,
                    "value axis must not be hidden",
                )
                self.assertIsNotNone(chart.x_axis.title)
                self.assertIsNotNone(chart.y_axis.title)

    def _seed_configured_pair(self, database: FluxDatabase) -> None:
        """One covered subscription, one configured gap subscription."""
        seed_month(database, date(2026, 7, 1), 1000.0)
        database.save_integration({
            "name": "Azure", "tenantId": "tenant", "enabled": True,
            "authMode": "managed_identity",
            "subscriptions": [
                {"subscriptionId": "sub-a", "label": "Covered"},
                {"subscriptionId": "sub-gap", "label": "Gap"},
            ],
        })
        database.save_budget_groups([
            {"name": "Everything", "annualAmount": 9000.0,
             "currency": "USD", "subscriptionIds": ["sub-a", "sub-gap"]},
        ])

    @staticmethod
    def _seed_daily(database: FluxDatabase, sub: str, service: str,
                    amount: float, day: date | None = None) -> None:
        from datetime import timedelta

        today = date.today()
        if day is None:
            day = max(today.replace(day=1), today - timedelta(days=3))
        with database.connect() as db:
            db.execute(
                "INSERT OR REPLACE INTO daily_cost_history VALUES "
                "(?,?,?,?,?,?,?,?,?,?)",
                ["test-daily", datetime.now(timezone.utc), day,
                 "ActualCost", sub, f"/subscriptions/{sub}/rg/x",
                 service, amount, "USD", "azure_focus_export"],
            )
            db.commit()

    @staticmethod
    def _cells(workbook, sheet_name):
        return [
            str(cell.value)
            for row in workbook[sheet_name].iter_rows()
            for cell in row
            if cell.value is not None
        ]

    def test_workbook_is_business_facing(self):
        """2026-08-11 feedback: warnings out, measured impact in.

        The workbook must not carry the operational coverage warning or
        the follow-up subscription table, must state a material gap's
        measured spend instead, must fold budget groups into the FY
        outlook sheet, and must present the spend composition on the
        Summary sheet as fiscal-year actuals split by budget group.
        """
        from datetime import timedelta

        from api.executive_export import build_executive_workbook

        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            self._seed_configured_pair(database)
            # A failed collection (admin-worthy in the app) with real
            # measured spend: material to a ~$1,000/month projection.
            database.start_cost_history_run("run-x", 1)
            database.begin_cost_history_scope(
                "run-x", "sub-gap", "AmortizedCost",
                date(2026, 7, 1), date(2026, 8, 5),
            )
            database.finish_cost_history_scope(
                "run-x", "sub-gap", "AmortizedCost", status="failed",
                status_code=429, message="HTTP 429", retained_last_good=False,
            )
            self._seed_daily(database, "sub-gap", "Virtual Machines", 50.0)
            self._seed_daily(database, "sub-a", "Virtual Machines", 40.0)
            self._seed_daily(database, "sub-a", "Storage", 10.0)
            # A finalized day in the previous month, so the FY-actuals
            # composition has data whether or not the fiscal year already
            # has a complete month behind it.
            previous_month = date.today().replace(day=1) - timedelta(days=15)
            self._seed_daily(
                database, "sub-a", "Virtual Machines", 40.0,
                day=previous_month,
            )

            outlook = database.fiscal_year_outlook()
            self.assertTrue(
                any("administrator" in item for item in outlook["limitations"]),
                "the app keeps the operational warning",
            )

            workbook = load_workbook(
                BytesIO(build_executive_workbook(database))
            )
            self.assertNotIn("Planning lens", workbook.sheetnames)
            self.assertNotIn("Budget groups", workbook.sheetnames)
            self.assertNotIn("Service composition", workbook.sheetnames)

            summary = " | ".join(self._cells(workbook, "Summary"))
            self.assertNotIn("administrator attention", summary)
            self.assertNotIn(
                "Subscriptions requiring cost-history follow-up", summary
            )
            self.assertIn("The outlook excludes 1 subscription(s)", summary)
            self.assertIn("measured spend is ~$", summary)
            # The composition table: FY actuals, grouped columns, totals.
            # (On the fiscal year's very first day there are no finalized
            # FY days yet and the table is legitimately absent.)
            today = date.today()
            if today != date(today.year, 7, 1):
                self.assertIn("Spend composition", summary)
                self.assertIn("Category", summary)
                self.assertIn("Everything", summary)
                self.assertIn("Virtual machine compute", summary)
            self.assertNotIn("Billing-service view", summary)
            self.assertNotIn("Cost source", summary)
            self.assertNotIn("billing classified", summary)

            fy_outlook = " | ".join(self._cells(workbook, "FY outlook"))
            self.assertIn("Budget groups", fy_outlook)
            self.assertIn("Everything", fy_outlook)
            self.assertNotIn("Subscriptions backfilled", fy_outlook)
            self.assertIn("excludes ~$", fy_outlook)

    def test_immaterial_gaps_are_omitted_entirely(self):
        from api.executive_export import build_executive_workbook

        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            self._seed_configured_pair(database)
            # sub-gap has no history and no measured spend at all.
            workbook = load_workbook(
                BytesIO(build_executive_workbook(database))
            )
            summary = " | ".join(self._cells(workbook, "Summary"))
            self.assertNotIn("administrator attention", summary)
            self.assertNotIn("The outlook excludes", summary)
            fy_outlook = " | ".join(self._cells(workbook, "FY outlook"))
            self.assertNotIn("excludes ~$", fy_outlook)

    def test_percent_cells_scale_instead_of_appending_a_literal_sign(self):
        """A literal '0.0"%"' rendered a group running 49.5% over as 0.5%.

        Both percent call sites pass fractions (variance/annualBudget and
        percentOfTotal/100), so the format must be one Excel scales by 100.
        A literal-suffix format loses two orders of magnitude silently --
        the number still looks plausible, which is why it shipped.
        """
        from api.executive_export import build_executive_workbook

        with TemporaryDirectory() as tmp:
            database = make_database(tmp)
            seed_month(database, date(2026, 7, 1), 1000.0)
            workbook = load_workbook(
                BytesIO(build_executive_workbook(database))
            )
            formats = {
                cell.number_format
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.number_format and "%" in cell.number_format
            }
        for number_format in formats:
            self.assertNotIn(
                '"%"', number_format,
                "a literal percent suffix does not scale the fraction",
            )


if __name__ == "__main__":
    unittest.main()
